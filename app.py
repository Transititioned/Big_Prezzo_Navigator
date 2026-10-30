"""Local-first Gradio interface for the Big Prezo PowerPoint engine."""
from __future__ import annotations

import json
import logging
import logging.handlers
import tempfile
import time
import uuid
from pathlib import Path
from typing import Any

import gradio as gr
from pptx import Presentation
from pptx.exc import PackageNotFoundError

from big_prezo import manifest_for_slides, process_presentation

MAX_UPLOAD_BYTES = 250 * 1024 * 1024
OUTPUT_DIR = Path(tempfile.gettempdir()) / "big_prezo_outputs"
OUTPUT_RETENTION_SECONDS = 24 * 60 * 60
LOG_PATH = OUTPUT_DIR / "big_prezo.log"

logger = logging.getLogger("big_prezo")
logger.setLevel(logging.INFO)
logger.propagate = False
if not logger.handlers:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    handler = logging.handlers.RotatingFileHandler(
        LOG_PATH, maxBytes=1_000_000, backupCount=3, encoding="utf-8"
    )
    handler.setFormatter(logging.Formatter(
        "%(asctime)s %(levelname)s %(name)s %(message)s"
    ))
    logger.addHandler(handler)
logger.info("Application started")


class UserFacingError(Exception):
    """A safe message to show without exposing exception details."""


def _clean_old_outputs() -> None:
    cutoff = time.time() - OUTPUT_RETENTION_SECONDS
    if not OUTPUT_DIR.is_dir():
        return
    for path in OUTPUT_DIR.iterdir():
        if path.is_file() and path.name.startswith("big_prezo_") and path.suffix in {".pptx", ".json"}:
            try:
                if path.stat().st_mtime < cutoff:
                    path.unlink(missing_ok=True)
            except OSError:
                logger.warning("Could not expire an old generated file")


def _validate_upload(upload: str | None, job_id: str) -> Path:
    if not upload:
        raise UserFacingError("Choose a PowerPoint file, then try again.")
    source = Path(upload)
    if source.suffix.lower() != ".pptx":
        raise UserFacingError("Choose a .pptx PowerPoint file.")
    if not source.is_file():
        raise UserFacingError("The uploaded file is no longer available. Upload it again.")
    try:
        size = source.stat().st_size
    except OSError:
        raise UserFacingError("The uploaded file could not be accessed. Upload it again.") from None
    if size > MAX_UPLOAD_BYTES:
        raise UserFacingError("This file is larger than the 250 MB upload limit.")
    if size == 0:
        raise UserFacingError("This PowerPoint file is empty or invalid. Choose another file.")
    return source


def _rows_from_table(table: Any) -> list[list[Any]]:
    if table is None:
        return []
    if hasattr(table, "tolist"):
        table = table.tolist()
    if hasattr(table, "values") and not isinstance(table, (list, tuple)):
        table = table.values.tolist()
    return [list(row) for row in table if row is not None]


def _safe_error(job_id: str, message: str) -> gr.Error:
    return gr.Error(f"{message} Reference: {job_id}")


def review_presentation(upload: str | None):
    """Populate the optional advanced section review controls."""
    job_id = uuid.uuid4().hex[:8]
    logger.info("job=%s stage=section_review_started", job_id)
    try:
        path = _validate_upload(upload, job_id)
        manifest = manifest_for_slides(Presentation(str(path)).slides)
    except UserFacingError as exc:
        raise _safe_error(job_id, str(exc)) from None
    except Exception:
        logger.exception("job=%s stage=section_review_failed", job_id)
        raise _safe_error(job_id, "We could not read this PowerPoint. Check that it opens correctly.") from None
    preview = [[r["slide_number"], r["title"] or "(Untitled slide)", r["compact_text"]]
               for r in manifest["slides"]]
    sections = [[r["slide_number"], r["title"] or "(Untitled slide)",
                 bool(r["probable_section_boundary"]), r["section"]]
                for r in manifest["slides"]]
    logger.info("job=%s stage=section_review_completed slides=%d", job_id, len(preview))
    return preview, sections, f"Optional review ready for {len(preview)} slides."


def generate_outputs(upload: str | None, table: Any = None, purpose: str = "",
                     takeaway: str = "", audience: str = ""):
    """Generate a navigable copy; never modify the uploaded source."""
    job_id = uuid.uuid4().hex[:8]
    logger.info("job=%s stage=generation_started", job_id)
    pptx_path: Path | None = None
    manifest_path: Path | None = None
    try:
        source = _validate_upload(upload, job_id)
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        _clean_old_outputs()
        token = uuid.uuid4().hex
        pptx_path = OUTPUT_DIR / f"big_prezo_{token}.pptx"
        manifest_path = OUTPUT_DIR / f"big_prezo_{token}.json"
        logger.info("job=%s stage=analysis_started", job_id)
        # Let python-pptx identify malformed packages with a nontechnical message.
        try:
            Presentation(str(source))
        except (PackageNotFoundError, KeyError, ValueError, OSError):
            raise UserFacingError("This PowerPoint file is invalid or damaged. Try opening and saving it in PowerPoint first.") from None
        logger.info("job=%s stage=navigation_generation_started", job_id)
        manifest = process_presentation(
            source, pptx_path, manifest_path,
            section_rows=_rows_from_table(table) if table is not None else None,
        )
        manifest["presentation_context"] = {
            "purpose": str(purpose or "").strip(),
            "main_takeaway": str(takeaway or "").strip(),
            "audience": str(audience or "").strip(),
        }
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        logger.info("job=%s stage=generation_completed slides=%d", job_id, manifest["source_slide_count"])
        # Gradio's File postprocessor requires a string path; returning Path
        # appears to succeed in the callback but produces an "Error" panel.
        return str(pptx_path), f"Your navigable PowerPoint is ready to download. Reference: {job_id}"
    except UserFacingError as exc:
        logger.warning("job=%s stage=generation_rejected reason=%s", job_id, str(exc))
        raise _safe_error(job_id, str(exc)) from None
    except Exception:
        logger.exception("job=%s stage=generation_failed", job_id)
        for candidate in (pptx_path, manifest_path):
            if candidate is not None:
                try:
                    candidate.unlink(missing_ok=True)
                except OSError:
                    logger.warning("job=%s stage=partial_output_cleanup_failed", job_id)
        raise _safe_error(job_id, "We could not create the navigable PowerPoint. Try another file or retry.") from None


CSS = """
.gradio-container { max-width: 980px !important; margin: auto !important; }
.hero { padding: 24px 28px; border-radius: 18px; background: linear-gradient(120deg,#102b46,#23618a); color: white; margin-bottom: 18px; }
.hero h1 { margin: 0 0 5px; font-size: 2rem; }
.hero p { margin: 0; opacity: .9; }
"""


def build_app() -> gr.Blocks:
    with gr.Blocks() as demo:
        gr.HTML("<div class='hero'><h1>Big Prezo</h1><p>Keep your slides. Find your way around.</p></div>")
        gr.Markdown("Upload a PowerPoint to create a new copy with clickable contents and Home links. Your file stays on this computer.")
        upload = gr.File(label="1. Upload PowerPoint", file_types=[".pptx"], type="filepath")
        with gr.Row():
            purpose = gr.Textbox(label="Purpose (optional)", placeholder="What is this presentation for?", scale=1)
            audience = gr.Textbox(label="Audience (optional)", placeholder="Who is it for?", scale=1)
        takeaway = gr.Textbox(label="Main takeaway (optional)", placeholder="What should people remember?", lines=1)
        with gr.Accordion("Advanced options", open=False):
            gr.Markdown("Review and adjust the automatically prepared contents links if you want.")
            preview_button = gr.Button("Review contents links", variant="secondary")
            advanced_status = gr.Markdown()
            section_table = gr.Dataframe(
                headers=["Slide", "Slide title", "Start section", "Section name"],
                datatype=["number", "str", "bool", "str"], type="array", interactive=True,
                max_height=420, label="Optional contents grouping",
            )
            preview_button.click(
                review_presentation, inputs=[upload], outputs=[gr.State(), section_table, advanced_status]
            )
        generate_button = gr.Button("2. Generate navigable PowerPoint", variant="primary")
        result_status = gr.Markdown()
        deck_output = gr.File(label="3. Download your PowerPoint")
        generate_button.click(
            generate_outputs,
            inputs=[upload, section_table, purpose, takeaway, audience],
            outputs=[deck_output, result_status],
            show_progress="full",
        )
        gr.Markdown("Your original presentation is never overwritten. The app supports PowerPoint files up to 250 MB.")
    return demo


demo = build_app()


if __name__ == "__main__":
    demo.launch(server_name="127.0.0.1", server_port=7860, inbrowser=True,
                share=False, theme=gr.themes.Soft(), css=CSS)
