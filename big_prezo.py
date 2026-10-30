from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any, Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

_MARKER_TOC = "BigPrezo::GeneratedToc"
_MARKER_HOME = "BigPrezo::Home::"
_SECTION_WORDS = re.compile(r"\b(part|section|chapter|appendix|appendices|workstream|agenda|overview|closing|q\s*&\s*a)\b", re.I)
_DEFAULT_ROWS_PER_TOC = 18


def _shape_text(shape: Any) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs).strip()


def _slide_text(slide: Any) -> list[str]:
    return [text for shape in slide.shapes if (text := _shape_text(shape))]


def _slide_title(slide: Any, texts: list[str]) -> str:
    title_shape = getattr(slide.shapes, "title", None)
    title = _shape_text(title_shape) if title_shape is not None else ""
    if title:
        return title
    return next((line.strip() for text in texts for line in text.splitlines() if line.strip()), "")


def probable_section_boundary(title: str, compact_text: str, *, max_body_chars: int = 140) -> bool:
    """Return true when a titled slide looks like a sparse section divider."""
    title = (title or "").strip()
    if not title:
        return False
    body = compact_text.replace(title, "", 1).strip()
    explicit = bool(_SECTION_WORDS.search(title))
    sparse = len(body) <= max_body_chars
    # Many real-world decks use plain topic titles (for example, "Email
    # Templates") for divider slides. Sparse content is a strong structural
    # signal even when the title has no conventional section keyword.
    return sparse and (explicit or len(body) <= 48)


def manifest_for_slides(slides: Iterable[Any], *, compact_limit: int = 500) -> dict[str, Any]:
    records = []
    section = "Opening"
    for number, slide in enumerate(slides, 1):
        texts = _slide_text(slide)
        title = _slide_title(slide, texts)
        compact = re.sub(r"\s+", " ", " | ".join(texts)).strip()[:compact_limit]
        boundary = probable_section_boundary(title, compact)
        if boundary:
            section = title
        records.append({
            "slide_number": number,
            "title": title,
            "compact_text": compact,
            "probable_section_boundary": boundary,
            "section": section,
        })
    section_count = sum(record["probable_section_boundary"] for record in records)
    return {
        "source_slide_count": len(records),
        "section_detection": {
            "probable_section_count": section_count,
            "all_slides_fallback": section_count == 0,
        },
        "slides": records,
    }


def apply_section_overrides(manifest: dict[str, Any], rows: Iterable[Iterable[Any]]) -> dict[str, Any]:
    """Apply reviewed per-slide section starts and names to a manifest.

    Rows are (slide number, title, start section, section name). A section name
    is taken from a row only when that row starts a section; other rows inherit.
    """
    by_number: dict[int, tuple[bool, str]] = {}
    for row in rows:
        if len(row) < 4:
            continue
        try:
            number = int(row[0])
        except (TypeError, ValueError):
            continue
        if number < 1 or number > manifest["source_slide_count"]:
            continue
        boundary = row[2] is True or str(row[2]).strip().lower() in {"true", "yes", "1", "y"}
        by_number[number] = (boundary, str(row[3] or "").strip())
    section = "Opening"
    for record in manifest["slides"]:
        boundary, name = by_number.get(record["slide_number"], (record["probable_section_boundary"], record["section"]))
        if boundary:
            section = name or f"Section at slide {record['slide_number']}"
        record["probable_section_boundary"] = boundary
        record["section"] = section
    count = sum(bool(record["probable_section_boundary"]) for record in manifest["slides"])
    manifest["section_detection"] = {
        "probable_section_count": count,
        "all_slides_fallback": count == 0,
    }
    return manifest


def _remove_shape(slide: Any, shape: Any) -> None:
    shape._element.getparent().remove(shape._element)


def _delete_slide(prs: Any, slide: Any) -> None:
    for sld_id in list(prs.slides._sldIdLst):
        if prs.part.related_part(sld_id.rId).partname == slide.part.partname:
            prs.slides._sldIdLst.remove(sld_id)
            prs.part.drop_rel(sld_id.rId)
            return


def _remove_generated_navigation(prs: Any) -> None:
    """Remove only shapes/slides explicitly marked as engine-owned."""
    for slide in list(prs.slides):
        generated_toc = any(shape.name == _MARKER_TOC for shape in slide.shapes)
        if generated_toc:
            _delete_slide(prs, slide)
            continue
        for shape in list(slide.shapes):
            if shape.name.startswith(_MARKER_HOME):
                _remove_shape(slide, shape)


def _set_hyperlink(shape: Any, target_slide: Any) -> None:
    shape.click_action.target_slide = target_slide


def _add_text_shape(slide: Any, x: int, y: int, w: int, h: int, text: str,
                    *, font_size: int, color: RGBColor, bold: bool = False,
                    name: str | None = None, align: Any = PP_ALIGN.LEFT,
                    fill: RGBColor | None = None) -> Any:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if fill else MSO_SHAPE.RECTANGLE,
                                   x, y, w, h)
    if name:
        shape.name = name
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.fill.background()
    else:
        shape.fill.background()
        shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    if name and name.startswith("BigPrezo::TocEntry::"):
        # Keep a visible cue in edit mode as well as the shape-level jump used
        # during Slide Show. The shape action remains the primary click target.
        run.font.color.rgb = RGBColor(31, 78, 121)
        run.font.underline = True
    return shape


def _insert_slides_at_front(prs: Any, new_slides: list[Any]) -> None:
    sld_ids = prs.slides._sldIdLst
    for slide in reversed(new_slides):
        for sld_id in list(sld_ids):
            if prs.part.related_part(sld_id.rId).partname == slide.part.partname:
                sld_ids.remove(sld_id)
                sld_ids.insert(0, sld_id)
                break


def _rgb_from_title(slides: list[Any]) -> RGBColor:
    """Reuse the source deck's first usable title colour, with a restrained fallback."""
    for slide in slides:
        title = getattr(slide.shapes, "title", None)
        if title is None or not title.has_text_frame:
            continue
        for paragraph in title.text_frame.paragraphs:
            for run in paragraph.runs:
                try:
                    color = run.font.color.rgb
                except (AttributeError, TypeError):
                    color = None
                if color is not None:
                    return RGBColor(color[0], color[1], color[2])
    return RGBColor(45, 71, 90)


def _home_button_position(slides: list[Any], width: int, height: int) -> tuple[int, int]:
    """Choose one low-overlap footer position shared by the source slides."""
    button_w, button_h = Inches(1.02), Inches(0.34)
    margin, gap = Inches(0.22), Inches(0.12)
    right = width - button_w - margin
    bottom = height - button_h - margin
    # Check common alignments and footer-object edges to find gaps between logos/text.
    x_candidates = {margin, right, (width - button_w) // 2}
    footer_top = height - Inches(0.85)
    for slide in slides:
        for shape in slide.shapes:
            if shape.name.startswith(_MARKER_HOME) or shape.top + shape.height < footer_top:
                continue
            x_candidates.add(max(margin, min(right, shape.left - button_w - gap)))
            x_candidates.add(max(margin, min(right, shape.left + shape.width + gap)))
    # Prefer the footer row; use one row above it only when every footer option is worse.
    y_candidates = [bottom, bottom - Inches(0.46)]
    positions = [(x, y) for y in y_candidates for x in sorted(x_candidates)]

    def score(pos: tuple[int, int]) -> tuple[int, int, int, int]:
        x, y = pos
        total = worst = 0
        for slide in slides:
            overlap = 0
            for shape in slide.shapes:
                if shape.name.startswith(_MARKER_HOME):
                    continue
                overlap_w = max(0, min(x + button_w, shape.left + shape.width) - max(x, shape.left))
                overlap_h = max(0, min(y + button_h, shape.top + shape.height) - max(y, shape.top))
                overlap += overlap_w * overlap_h
            total += overlap
            worst = max(worst, overlap)
        return total, worst, y_candidates.index(y), x

    return min(positions, key=score)


def _add_home_button(slide: Any, target: Any, number: int, width: int, height: int,
                     accent: RGBColor, position: tuple[int, int] | None = None) -> None:
    button_w, button_h = Inches(1.02), Inches(0.34)
    x, y = position or _home_button_position([slide], width, height)
    button = _add_text_shape(slide, x, y, button_w, button_h, "HOME · TOC",
                             font_size=9, color=RGBColor(255, 255, 255), bold=True,
                             name=f"{_MARKER_HOME}{number}", align=PP_ALIGN.CENTER,
                             fill=accent)
    _set_hyperlink(button, target)


def _clean_title(value: Any) -> str:
    value = str(value or "")
    value = value.replace("_x000B_", " ").replace("_x000b_", " ")
    value = re.sub(r"[\x00-\x1f\x7f]", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def _section_label(record: dict[str, Any]) -> str:
    return _clean_title(record.get("section") or "Slides")


def _toc_items(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Flatten groups into section headings and entries, suppressing divider duplicates."""
    items: list[dict[str, Any]] = []
    last_section = None
    for record in records:
        section = _section_label(record)
        title = _clean_title(record.get("title")) or "(Untitled slide)"
        boundary = bool(record.get("probable_section_boundary"))
        same_title = title.casefold() == section.casefold()
        if section != "Opening" and (section != last_section or boundary):
            items.append({"kind": "section", "text": section,
                          "slide_number": record["slide_number"]})
        if not (boundary and same_title):
            items.append({"kind": "entry", "text": title,
                          "slide_number": record["slide_number"]})
        last_section = section
    return items


def _add_toc_slides(prs: Any, source_slides: list[Any], manifest: dict[str, Any],
                    *, rows_per_toc: int = _DEFAULT_ROWS_PER_TOC) -> list[Any]:
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
    records = manifest["slides"]
    concise = len(records) >= 30 and any(record["probable_section_boundary"] for record in records)
    if concise:
        grouped: list[tuple[str, list[dict[str, Any]]]] = []
        for record in records:
            label = _section_label(record)
            if not grouped or grouped[-1][0] != label:
                grouped.append((label, []))
            grouped[-1][1].append(record)
        toc_records = []
        for label, group in grouped:
            first, last = group[0]["slide_number"], group[-1]["slide_number"]
            title = f"{label} · slides {first}–{last}" if first != last else f"{label} · slide {first}"
            toc_records.append({"slide_number": first, "title": title, "section": label,
                                "probable_section_boundary": True})
    else:
        toc_records = records
    if concise:
        items = [{"kind": "section", "text": _clean_title(record["title"]),
                  "slide_number": record["slide_number"]} for record in toc_records]
    else:
        items = _toc_items(toc_records)
    page_capacity = max(2, rows_per_toc)
    chunks = [items[i:i + page_capacity] for i in range(0, len(items), page_capacity)] or [[]]
    toc_slides: list[Any] = [prs.slides.add_slide(layout) for _ in chunks]
    width, height = prs.slide_width, prs.slide_height
    accent = _rgb_from_title(source_slides)
    for page_index, chunk in enumerate(chunks):
        slide = toc_slides[page_index]
        # Layout 6 is usually blank, but custom masters can still expose title
        # or content placeholders. Remove placeholders from engine-owned pages.
        for shape in list(slide.shapes):
            if getattr(shape, "is_placeholder", False):
                _remove_shape(slide, shape)
        marker = _add_text_shape(slide, 0, 0, 1, 1, "", font_size=1,
                                 color=accent, name=_MARKER_TOC)
        marker.left = -Inches(1)
        marker.top = -Inches(1)
        heading = "Contents" if len(chunks) == 1 else f"Contents · {page_index + 1} of {len(chunks)}"
        _add_text_shape(slide, Inches(0.58), Inches(0.30), width - Inches(1.16), Inches(0.48),
                        heading, font_size=25, color=accent, bold=True,
                        name="BigPrezo::TocHeading")
        top, bottom = Inches(1.08), Inches(0.70)
        # Keep small contents lists visually compact; use a second column once
        # a single column would become tall enough to feel sparse.
        columns = 1 if len(chunk) <= 7 else 2
        per_column = (len(chunk) + columns - 1) // columns
        gap = Inches(0.34)
        side = Inches(0.58)
        col_w = (width - side * 2 - gap) // columns
        available_h = height - top - bottom
        row_h = min(Inches(0.58), available_h // max(1, per_column))
        # Center short lists in the usable region without stretching their rows.
        list_h = row_h * per_column
        list_top = top + max(0, (available_h - list_h) // 2)
        for item_index, item in enumerate(chunk):
            col, row_index = divmod(item_index, per_column)
            row_y = list_top + row_index * row_h
            x = side + col * (col_w + gap)
            is_section = item["kind"] == "section"
            color = accent if is_section else RGBColor(42, 50, 58)
            shape = _add_text_shape(
                slide, x, row_y, col_w, row_h - Inches(0.035), item["text"],
                font_size=12 if is_section else 11, color=color, bold=is_section,
                name=f"BigPrezo::TocEntry::{item['slide_number']}")
            shape.text_frame.word_wrap = True
            if is_section and shape.text_frame.paragraphs[0].runs:
                # Keep linked section headings in the deck's accent colour;
                # ordinary entries retain their quieter hyperlink styling.
                shape.text_frame.paragraphs[0].runs[0].font.color.rgb = accent
                shape.text_frame.paragraphs[0].runs[0].font.underline = False
            target = source_slides[item["slide_number"] - 1]
            # Section headings can themselves represent linked divider slides.
            _set_hyperlink(shape, target)
        if page_index > 0:
            prev = _add_text_shape(slide, Inches(0.58), height - Inches(0.55), Inches(2.1), Inches(0.3),
                                   "‹ Previous contents", font_size=10, color=accent,
                                   name="BigPrezo::TocPrevious")
            _set_hyperlink(prev, toc_slides[page_index - 1])
        if page_index + 1 < len(chunks):
            nxt = _add_text_shape(slide, width - Inches(2.68), height - Inches(0.55), Inches(2.1), Inches(0.3),
                                  "Next contents ›", font_size=10, color=accent,
                                  align=PP_ALIGN.RIGHT, name="BigPrezo::TocNext")
            _set_hyperlink(nxt, toc_slides[page_index + 1])
    return toc_slides

def process_presentation(input_path: str | Path, output_path: str | Path,
                         manifest_path: str | Path | None = None, *,
                         rows_per_toc: int = _DEFAULT_ROWS_PER_TOC,
                         section_rows: Iterable[Iterable[Any]] | None = None) -> dict[str, Any]:
    if type(rows_per_toc) is not int or rows_per_toc <= 0:
        raise ValueError(f"rows_per_toc must be a positive integer; got {rows_per_toc!r}.")
    source = Path(input_path).resolve()
    output = Path(output_path).resolve()
    if source == output:
        raise ValueError("Output must be a new path; the original PPTX is never overwritten.")
    if not source.is_file():
        raise FileNotFoundError(source)
    prs = Presentation(str(source))
    # Keep old engine TOCs in the slide list until replacements are allocated.
    # python-pptx derives the next part name from slide count; deleting first can
    # reuse a live source slide's part name and corrupt repeated processing.
    old_toc_slides = [slide for slide in prs.slides
                      if any(shape.name == _MARKER_TOC for shape in slide.shapes)]
    original_slides = [slide for slide in prs.slides if slide not in old_toc_slides]
    for slide in original_slides:
        for shape in list(slide.shapes):
            if shape.name.startswith(_MARKER_HOME):
                _remove_shape(slide, shape)
    manifest = manifest_for_slides(original_slides)
    if section_rows is not None:
        apply_section_overrides(manifest, section_rows)
    toc_slides = _add_toc_slides(prs, original_slides, manifest, rows_per_toc=rows_per_toc)
    for old_slide in old_toc_slides:
        _delete_slide(prs, old_slide)
    home_position = _home_button_position(original_slides, prs.slide_width, prs.slide_height)
    for number, slide in enumerate(original_slides, 1):
        _add_home_button(slide, toc_slides[0], number, prs.slide_width, prs.slide_height,
                         _rgb_from_title(original_slides), home_position)
    _insert_slides_at_front(prs, toc_slides)
    toc_count = len(toc_slides)
    manifest["toc_slide_count"] = toc_count
    manifest["contents_slide_numbers"] = list(range(1, toc_count + 1))
    for record in manifest["slides"]:
        record["output_slide_number"] = record["slide_number"] + toc_count
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    if manifest_path:
        path = Path(manifest_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def extract_manifest_file(input_path: str | Path, output_path: str | Path) -> dict[str, Any]:
    prs = Presentation(str(input_path))
    _remove_generated_navigation(prs)
    manifest = manifest_for_slides(prs.slides)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Create a deterministic, navigable PowerPoint copy.")
    sub = parser.add_subparsers(dest="command", required=True)
    build = sub.add_parser("build", help="Add linked contents and Home controls to a copy of a PPTX.")
    build.add_argument("input", type=Path)
    build.add_argument("output", type=Path)
    build.add_argument("--manifest", type=Path, help="Optional JSON manifest output path.")
    build.add_argument("--rows-per-toc", type=int, default=_DEFAULT_ROWS_PER_TOC)
    manifest_cmd = sub.add_parser("manifest", help="Extract a JSON manifest without modifying the deck.")
    manifest_cmd.add_argument("input", type=Path)
    manifest_cmd.add_argument("output", type=Path)
    args = parser.parse_args(argv)
    if args.command == "build":
        result = process_presentation(args.input, args.output, args.manifest, rows_per_toc=args.rows_per_toc)
        print(f"Created {args.output} ({result['source_slide_count']} source slides; {result['toc_slide_count']} contents slides).")
    else:
        result = extract_manifest_file(args.input, args.output)
        print(f"Wrote {args.output} ({result['source_slide_count']} slides).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
