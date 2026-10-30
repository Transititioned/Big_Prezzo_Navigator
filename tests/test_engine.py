import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from big_prezo import apply_section_overrides, extract_manifest_file, process_presentation, probable_section_boundary


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / "source.pptx"
        self._make_deck(self.source)

    def tearDown(self):
        self.temp.cleanup()

    @staticmethod
    def _make_deck(path: Path):
        prs = Presentation()
        title_layout = prs.slide_layouts[5]
        blank = prs.slide_layouts[6]
        slide = prs.slides.add_slide(title_layout)
        slide.shapes.title.text = "Quarterly update"
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
        box.text = "Original opening content"
        slide.notes_slide.notes_text_frame.text = "Presenter note: preserve this exactly."

        slide = prs.slides.add_slide(title_layout)
        slide.shapes.title.text = "Part 1 — Delivery"
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
        box.text = "Section divider"

        slide = prs.slides.add_slide(blank)  # intentionally empty / titleless
        slide = prs.slides.add_slide(title_layout)
        slide.shapes.title.text = "Risks and next steps"
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
        box.text = "Detailed original slide text\nwith a second line"
        prs.save(path)

    def test_content_order_and_text_are_preserved(self):
        output = self.root / "navigable.pptx"
        manifest = process_presentation(self.source, output)
        result = Presentation(output)
        source = Presentation(self.source)
        toc_count = manifest["toc_slide_count"]
        self.assertEqual(len(result.slides), len(source.slides) + toc_count)
        for i, original in enumerate(source.slides):
            generated = result.slides[toc_count + i]
            self.assertEqual(
                [s.text for s in generated.shapes if s.has_text_frame and s.text.strip() and not s.name.startswith("BigPrezo::")],
                [s.text for s in original.shapes if s.has_text_frame and s.text.strip()],
            )
        self.assertEqual([r["title"] for r in manifest["slides"]],
                         ["Quarterly update", "Part 1 — Delivery", "", "Risks and next steps"])
        self.assertEqual([r["output_slide_number"] for r in manifest["slides"]],
                         [toc_count + n for n in range(1, 5)])
        self.assertEqual(result.slides[toc_count].notes_slide.notes_text_frame.text,
                         source.slides[0].notes_slide.notes_text_frame.text)

    def test_toc_and_home_links_target_slide_objects(self):
        output = self.root / "linked.pptx"
        manifest = process_presentation(self.source, output)
        result = Presentation(output)
        toc_count = manifest["toc_slide_count"]
        toc_targets = []
        for toc_slide in list(result.slides)[:toc_count]:
            for shape in toc_slide.shapes:
                if shape.name.startswith("BigPrezo::TocEntry::"):
                    toc_targets.append(shape.click_action.target_slide.slide_id)
        content_ids = [slide.slide_id for slide in list(result.slides)[toc_count:]]
        self.assertEqual(toc_targets, content_ids)
        for i, slide in enumerate(list(result.slides)[toc_count:], 1):
            buttons = [s for s in slide.shapes if s.name == f"BigPrezo::Home::{i}"]
            self.assertEqual(len(buttons), 1)
            self.assertEqual(buttons[0].click_action.target_slide.slide_id,
                             result.slides[0].slide_id)

    def test_generated_contents_are_compact_placeholder_free_and_home_is_visible(self):
        output = self.root / "polished.pptx"
        manifest = process_presentation(self.source, output)
        result = Presentation(output)
        toc = result.slides[0]
        self.assertFalse([shape for shape in toc.shapes if getattr(shape, "is_placeholder", False)])
        entries = [shape for shape in toc.shapes if shape.name.startswith("BigPrezo::TocEntry::")]
        self.assertEqual(len(entries), 4)
        # A short list should use compact rows rather than stretching links to
        # fill almost the entire slide height.
        self.assertLess(max(shape.height for shape in entries), Inches(0.60))
        self.assertLess(max(shape.top for shape in entries) - min(shape.top for shape in entries), Inches(3.0))

        source = Presentation(self.source)
        for index, original in enumerate(source.slides, 1):
            generated = result.slides[manifest["toc_slide_count"] + index - 1]
            button = next(shape for shape in generated.shapes
                          if shape.name == f"BigPrezo::Home::{index}")
            self.assertTrue(button.fill.type is not None)
            self.assertGreater(button.width, Inches(0.9))
            self.assertEqual(button.click_action.target_slide.slide_id, toc.slide_id)
            # Original placeholders remain on source slides; cleanup applies
            # only to generated contents pages.
            self.assertEqual(sum(getattr(s, "is_placeholder", False) for s in generated.shapes),
                             sum(getattr(s, "is_placeholder", False) for s in original.shapes))

    def test_home_button_uses_shared_gap_between_footer_items(self):
        footer = self.root / "footer.pptx"
        prs = Presentation()
        for index in range(3):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            # Distinct footer arrangements leave the same clear centre gap.
            slide.shapes.add_textbox(Inches(0.25), Inches(7.0), Inches(1.8), Inches(0.3)).text = f"Logo {index}"
            slide.shapes.add_textbox(Inches(11.25), Inches(7.0), Inches(1.8), Inches(0.3)).text = f"Footer {index}"
            if index == 1:
                slide.shapes.add_textbox(Inches(5.1), Inches(7.0), Inches(3.0), Inches(0.3)).text = "Page number"
        prs.save(footer)
        output = self.root / "footer_nav.pptx"
        manifest = process_presentation(footer, output)
        result = Presentation(output)
        source_slides = list(result.slides)[manifest["toc_slide_count"]:]
        buttons = [next(shape for shape in slide.shapes if shape.name.startswith("BigPrezo::Home::"))
                   for slide in source_slides]
        self.assertEqual(len({(button.left, button.top) for button in buttons}), 1)
        for slide, button in zip(source_slides, buttons):
            for shape in slide.shapes:
                if shape is button or shape.name.startswith("BigPrezo::Home::"):
                    continue
                overlap_w = max(0, min(button.left + button.width, shape.left + shape.width)
                                - max(button.left, shape.left))
                overlap_h = max(0, min(button.top + button.height, shape.top + shape.height)
                                - max(button.top, shape.top))
                self.assertEqual(overlap_w * overlap_h, 0, shape.name)
            self.assertEqual(button.click_action.target_slide.slide_id, result.slides[0].slide_id)

    def test_contents_generation_with_title_and_content_layout(self):
        alternate = self.root / "alternate_layout.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Different layout"
        slide.placeholders[1].text = "Original content placeholder stays on the source slide."
        prs.save(alternate)
        output = self.root / "alternate_layout_nav.pptx"
        manifest = process_presentation(alternate, output)
        generated = Presentation(output)
        self.assertFalse([shape for shape in generated.slides[0].shapes
                          if getattr(shape, "is_placeholder", False)])
        source_slide = generated.slides[manifest["toc_slide_count"]]
        self.assertTrue(any(getattr(shape, "is_placeholder", False)
                            and shape.text == "Original content placeholder stays on the source slide."
                            for shape in source_slide.shapes))

    def test_section_detection_and_no_section_fallback(self):
        self.assertTrue(probable_section_boundary("Part 2 — Delivery", "Part 2 — Delivery | divider"))
        self.assertFalse(probable_section_boundary("Quarterly update", "Quarterly update | five bullets"))
        plain = self.root / "plain.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Welcome"
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Metrics"
        prs.save(plain)
        output = self.root / "plain_nav.pptx"
        manifest = process_presentation(plain, output)
        self.assertFalse(any(r["probable_section_boundary"] for r in manifest["slides"]))
        self.assertTrue(manifest["section_detection"]["all_slides_fallback"])
        linked_rows = [s for slide in list(Presentation(output).slides)[:manifest["toc_slide_count"]]
                       for s in slide.shapes if s.name.startswith("BigPrezo::TocEntry::")]
        self.assertEqual(len(linked_rows), 2)

    def test_manifest_command_writes_json_and_creates_parent(self):
        output = self.root / "nested" / "manifest.json"
        manifest = extract_manifest_file(self.source, output)
        self.assertTrue(output.is_file())
        self.assertEqual(manifest["source_slide_count"], 4)

    def test_extract_manifest_omits_generated_contents_and_home_shapes(self):
        generated = self.root / "generated.pptx"
        manifest_path = self.root / "generated_manifest.json"
        process_presentation(self.source, generated)

        manifest = extract_manifest_file(generated, manifest_path)

        self.assertTrue(manifest_path.is_file())
        self.assertEqual(manifest["source_slide_count"], 4)
        self.assertEqual([record["title"] for record in manifest["slides"]],
                         ["Quarterly update", "Part 1 — Delivery", "", "Risks and next steps"])
        extracted_text = " ".join(record["compact_text"] for record in manifest["slides"])
        self.assertNotIn("HOME / TOC", extracted_text)
        self.assertNotIn("Contents", extracted_text)

    def test_rows_per_toc_rejects_invalid_values(self):
        for value in (0, -1, 1.5, "3", True, None):
            with self.subTest(value=value):
                with self.assertRaisesRegex(ValueError, "rows_per_toc must be a positive integer"):
                    process_presentation(self.source, self.root / "invalid.pptx",
                                         rows_per_toc=value)

    def test_reprocessing_does_not_accumulate_navigation(self):
        first = self.root / "first.pptx"
        second = self.root / "second.pptx"
        first_manifest = process_presentation(self.source, first)
        second_manifest = process_presentation(first, second)
        self.assertEqual(second_manifest["source_slide_count"], 4)
        self.assertEqual(second_manifest["toc_slide_count"], first_manifest["toc_slide_count"])
        output = Presentation(second)
        self.assertEqual(len(output.slides), 4 + second_manifest["toc_slide_count"])
        home_shapes = [shape for slide in output.slides for shape in slide.shapes
                       if shape.name.startswith("BigPrezo::Home::")]
        self.assertEqual(len(home_shapes), 4)
        self.assertEqual(sum(any(s.name == "BigPrezo::GeneratedToc" for s in slide.shapes)
                             for slide in output.slides), second_manifest["toc_slide_count"])

    def test_multiple_contents_pages_link_to_each_other(self):
        output = self.root / "paged.pptx"
        manifest = process_presentation(self.source, output, rows_per_toc=2)
        result = Presentation(output)
        self.assertEqual(manifest["toc_slide_count"], 2)
        first, second = list(result.slides)[:2]
        next_link = next(s for s in first.shapes if s.name == "BigPrezo::TocNext")
        previous_link = next(s for s in second.shapes if s.name == "BigPrezo::TocPrevious")
        self.assertEqual(next_link.click_action.target_slide.slide_id, second.slide_id)
        self.assertEqual(previous_link.click_action.target_slide.slide_id, first.slide_id)

    def test_empty_and_title_only_slides_are_safe(self):
        odd = self.root / "odd.pptx"
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Appendix"
        prs.save(odd)
        output = self.root / "odd_nav.pptx"
        manifest = process_presentation(odd, output)
        self.assertEqual(manifest["source_slide_count"], 2)
        self.assertEqual(manifest["slides"][0]["title"], "")
        self.assertTrue(manifest["slides"][1]["probable_section_boundary"])
        self.assertEqual(len(Presentation(output).slides), 3)

    def test_reviewed_section_correction_changes_manifest_and_toc(self):
        output = self.root / "corrected.pptx"
        rows = [
            [1, "Quarterly update", False, "Opening"],
            [2, "Part 1 — Delivery", False, "Delivery"],
            [3, "", True, "Delivery"],
            [4, "Risks and next steps", False, "Delivery"],
        ]
        manifest = process_presentation(self.source, output, section_rows=rows)
        self.assertEqual(manifest["slides"][1]["section"], "Opening")
        self.assertTrue(manifest["slides"][2]["probable_section_boundary"])
        self.assertEqual(manifest["slides"][2]["section"], "Delivery")
        result = Presentation(output)
        toc_text = " ".join(shape.text for shape in result.slides[0].shapes if shape.has_text_frame)
        self.assertIn("Delivery", toc_text)
        toc_entry = next(s for s in result.slides[0].shapes if s.name == "BigPrezo::TocEntry::3")
        self.assertEqual(toc_entry.click_action.target_slide.slide_id,
                         result.slides[manifest["toc_slide_count"] + 2].slide_id)

    def test_existing_native_hyperlink_is_preserved(self):
        linked_source = self.root / "linked_source.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Source slide"
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
        box.text = "Existing link"
        box.click_action.hyperlink.address = "https://example.com/presentation-reference"
        target = prs.slides.add_slide(prs.slide_layouts[5])
        target.shapes.title.text = "Existing internal target"
        internal = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
        internal.text = "Jump within deck"
        internal.click_action.target_slide = target
        prs.save(linked_source)
        output = self.root / "linked_copy.pptx"
        process_presentation(linked_source, output)
        copied = Presentation(output)
        source_slide = copied.slides[1]
        link_shape = next(shape for shape in source_slide.shapes if shape.has_text_frame and shape.text == "Existing link")
        self.assertEqual(link_shape.click_action.hyperlink.address,
                         "https://example.com/presentation-reference")
        internal_shape = next(shape for shape in source_slide.shapes
                              if shape.has_text_frame and shape.text == "Jump within deck")
        self.assertEqual(internal_shape.click_action.target_slide.slide_id, copied.slides[2].slide_id)

    def test_large_deck_uses_section_summary_contents(self):
        large = self.root / "large.pptx"
        prs = Presentation()
        for number in range(1, 36):
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = ({1: "Part 1 — Opening", 20: "Part 2 — Delivery"}.get(number)
                                        or f"Slide {number}")
        prs.save(large)
        output = self.root / "large_nav.pptx"
        manifest = process_presentation(large, output)
        toc = list(Presentation(output).slides)[0]
        entries = [s for s in toc.shapes if s.name.startswith("BigPrezo::TocEntry::")]
        self.assertEqual(len(entries), 2)
        self.assertIn("Part 1", " ".join(s.text for s in entries))
        self.assertIn("Part 2", " ".join(s.text for s in entries))


if __name__ == "__main__":
    unittest.main()
