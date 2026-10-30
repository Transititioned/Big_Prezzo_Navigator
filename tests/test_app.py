import json
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

import app
from app import generate_outputs, review_presentation


class GradioWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.source = Path(self.temp.name) / "source.pptx"
        prs = Presentation()
        for title in ("Welcome", "Project plan", "Risks"):
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = title
        prs.slides[0].notes_slide.notes_text_frame.text = "Keep my notes."
        prs.save(self.source)
        self.original_bytes = self.source.read_bytes()

    def tearDown(self):
        self.temp.cleanup()

    def test_review_extracts_slides_and_editable_section_rows(self):
        preview, sections, status = review_presentation(str(self.source))
        self.assertEqual(len(preview), 3)
        self.assertEqual(preview[1][1], "Project plan")
        self.assertEqual(sections[0][2], False)
        self.assertIn("3 slides", status)

    def test_edited_boundary_and_name_are_in_generated_deck_and_manifest(self):
        _, sections, _ = review_presentation(str(self.source))
        sections[1][2] = True
        sections[1][3] = "Delivery plan"
        pptx_file, status = generate_outputs(
            str(self.source), sections, "Board update", "Agree next steps", "Executives")
        pptx_path = Path(pptx_file)
        self.assertIsInstance(pptx_file, str)
        manifest_path = Path(app.OUTPUT_DIR) / f"{pptx_path.stem}.json"
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            self.assertTrue(manifest["slides"][1]["probable_section_boundary"])
            self.assertEqual(manifest["slides"][1]["section"], "Delivery plan")
            self.assertEqual(manifest["presentation_context"]["audience"], "Executives")
            deck = Presentation(pptx_path)
            entry = next(shape for shape in deck.slides[0].shapes if shape.name == "BigPrezo::TocEntry::2")
            self.assertIn("Delivery plan", entry.text)
            self.assertEqual(entry.click_action.target_slide.slide_id,
                             deck.slides[manifest["toc_slide_count"] + 1].slide_id)
            self.assertEqual(deck.slides[manifest["toc_slide_count"]].notes_slide.notes_text_frame.text,
                             "Keep my notes.")
            self.assertIn("ready to download", status)
            self.assertIn("Reference:", status)
            self.assertEqual(self.source.read_bytes(), self.original_bytes)
        finally:
            pptx_path.unlink(missing_ok=True)
            manifest_path.unlink(missing_ok=True)

    def test_generation_without_advanced_review_creates_openable_deck(self):
        pptx_file, status = generate_outputs(str(self.source))
        self.assertIsInstance(pptx_file, str)
        path = Path(pptx_file)
        try:
            self.assertTrue(path.is_file())
            self.assertEqual(len(Presentation(path).slides), 4)
            self.assertIn("ready to download", status)
        finally:
            path.unlink(missing_ok=True)
            path.with_suffix(".json").unlink(missing_ok=True)

    def test_missing_upload_has_clear_message_and_job_reference(self):
        with self.assertRaises(app.gr.Error) as raised:
            generate_outputs(None)
        self.assertIn("Choose a PowerPoint file", str(raised.exception))
        self.assertRegex(str(raised.exception), r"Reference: [0-9a-f]{8}")

    def test_invalid_powerpoint_has_clear_message_and_job_reference(self):
        invalid = Path(self.temp.name) / "broken.pptx"
        invalid.write_bytes(b"not a presentation")
        with self.assertRaises(app.gr.Error) as raised:
            generate_outputs(str(invalid))
        self.assertIn("invalid or damaged", str(raised.exception))
        self.assertRegex(str(raised.exception), r"Reference: [0-9a-f]{8}")

    def test_generation_failure_logs_traceback_without_deck_text(self):
        from unittest.mock import patch

        with patch.object(app, "process_presentation", side_effect=RuntimeError("controlled failure")):
            with self.assertRaises(app.gr.Error) as raised:
                generate_outputs(str(self.source))
        self.assertIn("could not create", str(raised.exception))
        self.assertRegex(str(raised.exception), r"Reference: [0-9a-f]{8}")
        self.assertTrue(app.LOG_PATH.exists())
        log = app.LOG_PATH.read_text(encoding="utf-8")
        self.assertIn("Traceback (most recent call last)", log)
        self.assertIn("controlled failure", log)
        self.assertNotIn("Keep my notes.", log)

    def test_advanced_review_is_collapsed_and_default_flow_is_simplified(self):
        config = app.build_app().get_config_file()
        labels = [component.get("props", {}).get("label") for component in config["components"]]
        self.assertIn("3. Download your PowerPoint", labels)
        self.assertNotIn("JSON manifest", labels)
        self.assertNotIn("section", " ".join(str(label) for label in labels).lower())


if __name__ == "__main__":
    unittest.main()
